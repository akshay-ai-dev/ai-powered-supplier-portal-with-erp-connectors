"""Readable shipment summary of a REVIEWED draft, plus TXT / PDF / DOCX exports.

Everything is built from the reviewed JSON that review.finish_review returns, i.e. the
supplier's final edited values. The browser modal and all three downloads use the same
build_summary() output, so they always show the same values.

  * TXT  - plain text (UTF-8).
  * PDF  - small built-in writer (standard Helvetica font, no extra package). Characters
           outside Windows-1252 are printed as '?'.
  * DOCX - python-docx (already installed as a Docling dependency; pinned in requirements-lock.txt).

Local POC only: nothing here submits a shipment or posts to ERP.
"""

from __future__ import annotations

import io
import re
from datetime import date, datetime

from review import REVIEWED_STATUS

BANNER = "Reviewed draft — not submitted to ERP"
NOTICE = ("Local POC only: this reviewed draft has NOT been submitted as a shipment "
          "and has NOT been posted to ERP.")
NOT_PROVIDED = "Not provided"
EXPORT_FORMATS = {
    "txt": "text/plain; charset=utf-8",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
FIELD_LABELS = {
    "shipDate": "Ship date", "carrier": "Carrier", "trackingNumbers": "Tracking numbers",
    "shippedQuantity": "Shipped quantity", "unitOfMeasure": "Unit", "lotNumbers": "Lot numbers",
    "serialNumbers": "Serial numbers", "items": "Item lines", "poComparison": "PO comparison",
}
PO_STATUS = {
    "match": "Matches the PO",
    "over_shipped": "Over-shipped",
    "under_shipped": "Under-shipped",
    "cannot_compare": "Cannot compare",
}


# --------------------------------------------------------------------------- summary
def _num(value):
    if value is None:
        return None
    x = round(float(value), 6)
    return str(int(x)) if x == int(x) else str(x)


def _qty(qty, unit):
    if qty is None:
        return NOT_PROVIDED
    return f"{_num(qty)} {unit}" if unit else f"{_num(qty)} (no unit)"


def _date(iso):
    if not iso:
        return NOT_PROVIDED
    try:
        d = date.fromisoformat(iso)
    except ValueError:
        return iso
    return f"{iso} ({d.day} {d.strftime('%B %Y')})"


def _field_label(field):
    m = re.fullmatch(r"items\[(\d+)\]\.(\w+)", field or "")
    if m:
        return f"Line {int(m.group(1)) + 1} {FIELD_LABELS.get(m.group(2), m.group(2))}".replace(
            "quantityShipped", "shipped quantity")
    return FIELD_LABELS.get(field, field)


def _show(value):
    if value is None or value == [] or value == "":
        return "empty"
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    if isinstance(value, float):
        return _num(value)
    return str(value)


def _review_notes(reviewed):
    review = reviewed.get("review") or {}
    notes = []
    for c in review.get("corrections") or []:
        notes.append(f"Corrected {_field_label(c.get('field'))}: {_show(c.get('extracted'))} "
                     f"→ {_show(c.get('reviewed'))}")
    messages = {f"{i.get('field')}:{i.get('code')}": i.get("message")
                for i in reviewed.get("fieldIssues") or []}
    for key in review.get("acknowledgedIssues") or []:
        field = key.split(":", 1)[0]
        notes.append(f"Checked by supplier ({_field_label(field)}): {messages.get(key) or key}")
    for w in review.get("warnings") or []:
        notes.append(f"Note: {w.get('message')}")
    return notes


def _po_result(review):
    po, c = review.get("po"), review.get("poComparison")
    if not po:
        return None
    rows = [("PO quantity", _qty(po.get("quantity"), po.get("unit"))),
            ("PO item / part no.", po.get("item") or NOT_PROVIDED)]
    if not c:
        rows.append(("Result", "Not compared"))
        return rows
    rows.append(("Result", PO_STATUS.get(c.get("status"), c.get("status") or NOT_PROVIDED)))
    if c.get("status") != "cannot_compare":
        diff = c.get("difference")
        sign = "+" if diff is not None and diff > 0 else ""
        rows.append(("Shipped (compared)", _qty(c.get("shippedQuantity"), c.get("shippedUnit"))))
        rows.append(("Difference", f"{sign}{_num(diff)} {c.get('shippedUnit') or ''}".strip()
                     if diff is not None else NOT_PROVIDED))
    for msg in c.get("issues") or []:
        rows.append(("Detail", msg))
    return rows


def build_summary(reviewed):
    """Plain data (labels and display strings) describing the final edited shipment."""
    review = reviewed.get("review") or {}
    lots, serials = reviewed.get("lotNumbers") or [], reviewed.get("serialNumbers") or []
    items = []
    for idx, it in enumerate(reviewed.get("items") or []):
        item_lots = [l.get("lotNumber") for l in it.get("lots") or [] if l.get("lotNumber") in lots]
        item_serials = [s for s in it.get("serialNumbers") or [] if s in serials]
        items.append({
            "line": str(it.get("lineNumber") or idx + 1),
            "item": it.get("itemNumber") or NOT_PROVIDED,
            "description": it.get("description") or NOT_PROVIDED,
            "shipped": _qty(it.get("quantityShipped"), it.get("unitOfMeasure")),
            "lots": ", ".join(item_lots) or NOT_PROVIDED,
            "serials": ", ".join(item_serials) or NOT_PROVIDED,
        })
    reviewed_at = review.get("reviewedAt")
    try:
        reviewed_at = datetime.fromisoformat(reviewed_at).strftime("%Y-%m-%d %H:%M UTC") if reviewed_at else None
    except ValueError:
        pass
    return {
        "banner": BANNER,
        "notice": NOTICE,
        "sourceFile": reviewed.get("sourceFile") or NOT_PROVIDED,
        "reviewedAt": reviewed_at or NOT_PROVIDED,
        "shipment": [
            ["Ship date", _date(reviewed.get("shipDate"))],
            ["Carrier", reviewed.get("carrier") or NOT_PROVIDED],
            ["Shipped quantity", _qty(reviewed.get("shippedQuantity"), reviewed.get("unitOfMeasure"))
             if reviewed.get("shippedQuantity") is not None
             else ("No single total (see item lines)" if items else NOT_PROVIDED)],
            ["Unit", reviewed.get("unitOfMeasure") or NOT_PROVIDED],
        ],
        "lists": [
            ["Tracking numbers", list(reviewed.get("trackingNumbers") or [])],
            ["Lot numbers", list(lots)],
            ["Serial numbers", list(serials)],
        ],
        "items": items,
        "notes": _review_notes(reviewed),
        "po": _po_result(review),
    }


def summary_from_reviewed(reviewed):
    if not isinstance(reviewed, dict) or reviewed.get("status") != REVIEWED_STATUS:
        raise ValueError("Finish the review first; only a reviewed draft can be exported.")
    return build_summary(reviewed)


# --------------------------------------------------------------------------- TXT
def _text_lines(s):
    """Shared line layout for TXT and PDF: (style, text) with style in title/head/body/blank."""
    out = [("title", s["banner"].upper()), ("body", s["notice"]), ("blank", ""),
           ("body", f"Source document: {s['sourceFile']}"), ("body", f"Reviewed at: {s['reviewedAt']}"),
           ("blank", ""), ("head", "Shipment")]
    out += [("body", f"{label}: {value}") for label, value in s["shipment"]]
    for label, values in s["lists"]:
        out.append(("body", f"{label}: {', '.join(values) if values else NOT_PROVIDED}"))
    out += [("blank", ""), ("head", "Item lines")]
    if not s["items"]:
        out.append(("body", NOT_PROVIDED))
    for it in s["items"]:
        out.append(("body", f"Line {it['line']}: {it['item']} — {it['description']}"))
        out.append(("indent", f"Shipped: {it['shipped']}"))
        out.append(("indent", f"Lot numbers: {it['lots']}"))
        out.append(("indent", f"Serial numbers: {it['serials']}"))
    out += [("blank", ""), ("head", "PO result")]
    out += [("body", f"{k}: {v}") for k, v in s["po"]] if s["po"] else [("body", NOT_PROVIDED)]
    out += [("blank", ""), ("head", "Review notes")]
    out += [("body", f"- {n}") for n in s["notes"]] if s["notes"] else [("body", "None")]
    out += [("blank", ""), ("body", s["banner"] + ".")]
    return out


def render_txt(s):
    lines = []
    for style, text in _text_lines(s):
        if style == "head":
            lines += [text, "-" * len(text)]
        elif style == "indent":
            lines.append("    " + text)
        else:
            lines.append(text)
    return ("\r\n".join(lines) + "\r\n").encode("utf-8")


# --------------------------------------------------------------------------- PDF
# Helvetica advance widths (1/1000 em) for ASCII 32..126, from the standard AFM metrics.
_HELV = [278, 278, 355, 556, 556, 889, 667, 191, 333, 333, 389, 584, 278, 333, 278, 278,
         556, 556, 556, 556, 556, 556, 556, 556, 556, 556, 278, 278, 584, 584, 584, 556,
         1015, 667, 667, 722, 722, 667, 611, 778, 722, 278, 500, 667, 556, 833, 722, 778,
         667, 778, 722, 667, 611, 722, 667, 944, 667, 667, 611, 278, 278, 278, 469, 556,
         333, 556, 556, 500, 556, 556, 278, 556, 556, 222, 222, 500, 222, 833, 556, 556,
         556, 556, 333, 500, 278, 556, 500, 722, 500, 500, 500, 334, 260, 334, 584]
_PAGE_W, _PAGE_H, _MARGIN = 595, 842, 56
_STYLES = {"title": ("F2", 14, 22), "head": ("F2", 11.5, 20), "body": ("F1", 10, 14),
           "indent": ("F1", 10, 14), "blank": ("F1", 10, 8)}


def _width(text, size, bold):
    w = sum(_HELV[ord(ch) - 32] if 32 <= ord(ch) <= 126 else (1000 if ch == "—" else 600) for ch in text)
    return w * size / 1000 * (1.07 if bold else 1.0)


def _wrap(text, size, bold, max_w):
    words = []
    for word in text.split(" "):  # break single tokens wider than the line
        while len(word) > 1 and _width(word, size, bold) > max_w:
            cut = len(word) - 1
            while cut > 1 and _width(word[:cut], size, bold) > max_w:
                cut -= 1
            words.append(word[:cut])
            word = word[cut:]
        words.append(word)
    lines, cur = [], ""
    for word in words:
        trial = f"{cur} {word}" if cur else word
        if cur and _width(trial, size, bold) > max_w:
            lines.append(cur)
            cur = word
        else:
            cur = trial
    lines.append(cur)
    return lines


def _pdf_str(text):
    raw = text.replace("→", "->").encode("cp1252", errors="replace")
    return b"(" + raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)") + b")"


def render_pdf(s):
    pages, ops, y = [], [], _PAGE_H - _MARGIN
    for style, text in _text_lines(s):
        text = text.replace("→", "->")
        font, size, lead = _STYLES[style]
        x = _MARGIN + (18 if style == "indent" else 0)
        wrapped = [""] if style == "blank" else _wrap(text, size, font == "F2", _PAGE_W - _MARGIN - x)
        for i, line in enumerate(wrapped):
            if y - lead < _MARGIN:
                pages.append(ops)
                ops, y = [], _PAGE_H - _MARGIN
            y -= lead
            if line:
                ops.append(b"BT /%s %g Tf %g %g Td %s Tj ET" % (font.encode(), size, x + (12 if i else 0), y,
                                                               _pdf_str(line)))
            if style == "head" and i == len(wrapped) - 1:
                ops.append(b"0.6 w %g %g m %g %g l S" % (_MARGIN, y - 4, _PAGE_W - _MARGIN, y - 4))
    pages.append(ops)
    footer = _pdf_str(BANNER)
    n = len(pages)
    # Object numbers: 1 catalog, 2 pages, 3 Helvetica, 4 Helvetica-Bold, then (page, content) pairs.
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               b"<< /Type /Pages /Kids [%s] /Count %d >>" % (b" ".join(b"%d 0 R" % (5 + 2 * i) for i in range(n)), n),
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>"]
    for i, page_ops in enumerate(pages):
        stream = b"\n".join(page_ops + [b"BT /F1 8 Tf %d 30 Td %s Tj ET" % (_MARGIN, footer),
                                        b"BT /F1 8 Tf %d 30 Td (Page %d of %d) Tj ET" % (_PAGE_W - _MARGIN - 50, i + 1, n)])
        objects.append(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 %d %d] /Resources << /Font << /F1 3 0 R "
                       b"/F2 4 0 R >> >> /Contents %d 0 R >>" % (_PAGE_W, _PAGE_H, 6 + 2 * i))
        objects.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream))
    title = _pdf_str(f"{BANNER} - {s['sourceFile']}")
    objects.append(b"<< /Title %s /Producer (Packing-List Pre-Fill POC) >>" % title)
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for num, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n%s\nendobj\n" % (num, body))
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    out.write(b"".join(b"%010d 00000 n \n" % off for off in offsets))
    out.write(b"trailer\n<< /Size %d /Root 1 0 R /Info %d 0 R >>\nstartxref\n%d\n%%%%EOF\n"
              % (len(objects) + 1, len(objects), xref))
    return out.getvalue()


# --------------------------------------------------------------------------- DOCX
def render_docx(s):
    from docx import Document  # python-docx
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.core_properties.title = f"{BANNER} - {s['sourceFile']}"
    for section in doc.sections:
        section.header.paragraphs[0].text = BANNER
        section.footer.paragraphs[0].text = BANNER
    doc.add_heading(BANNER, level=1)
    notice = doc.add_paragraph().add_run(s["notice"])
    notice.bold = True
    notice.font.color.rgb = RGBColor(0x9A, 0x34, 0x12)
    doc.add_paragraph(f"Source document: {s['sourceFile']}\nReviewed at: {s['reviewedAt']}")

    def table(rows, header=None):
        t = doc.add_table(rows=0, cols=len(rows[0]) if rows else len(header))
        t.style = "Table Grid"
        if header:
            cells = t.add_row().cells
            for c, text in zip(cells, header):
                c.text = ""
                c.paragraphs[0].add_run(text).bold = True
        for row in rows:
            cells = t.add_row().cells
            for c, text in zip(cells, row):
                c.text = str(text)
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(10)
        return t

    doc.add_heading("Shipment", level=2)
    table([*s["shipment"], *[[label, ", ".join(v) if v else NOT_PROVIDED] for label, v in s["lists"]]])
    doc.add_heading("Item lines", level=2)
    if s["items"]:
        table([[it["line"], it["item"], it["description"], it["shipped"], it["lots"], it["serials"]]
               for it in s["items"]],
              header=["Line", "Item", "Description", "Shipped", "Lot numbers", "Serial numbers"])
    else:
        doc.add_paragraph(NOT_PROVIDED)
    doc.add_heading("PO result", level=2)
    if s["po"]:
        table(s["po"])
    else:
        doc.add_paragraph(NOT_PROVIDED)
    doc.add_heading("Review notes", level=2)
    for n in s["notes"] or ["None"]:
        doc.add_paragraph(n, style="List Bullet")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


RENDERERS = {"txt": render_txt, "pdf": render_pdf, "docx": render_docx}


def export(reviewed, fmt):
    """Return (bytes, content_type, filename). Raises ValueError for bad input."""
    if fmt not in RENDERERS:
        raise ValueError("Unknown export format. Use txt, pdf or docx.")
    s = summary_from_reviewed(reviewed)
    stem = re.sub(r"[^A-Za-z0-9._-]", "_", re.sub(r"\.pdf$", "", str(reviewed.get("sourceFile") or "draft"),
                                                  flags=re.I))[:80] or "draft"
    return RENDERERS[fmt](s), EXPORT_FORMATS[fmt], f"{stem}-reviewed-draft.{fmt}"
