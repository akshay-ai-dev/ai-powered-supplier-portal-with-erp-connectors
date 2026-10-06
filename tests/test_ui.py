"""Tests for the browser-UI backend: review logic (edited-value PO comparison, validation)
and the local HTTP API with a fake converter (no Docling needed).

    .venv\\Scripts\\python.exe -m unittest tests.test_ui -v
"""

import copy
import json
import sys
import threading
import time
import unittest
import urllib.error
import urllib.request
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import review  # noqa: E402
import summary_export  # noqa: E402
import ui_server  # noqa: E402

ONE_ITEM = {
    "sourceFile": "x.pdf", "shipDate": "2024-01-16", "carrier": "UPS", "trackingNumbers": [],
    "shippedQuantity": 42.0, "unitOfMeasure": "PR", "lotNumbers": ["2401"], "serialNumbers": [],
    "items": [{"itemNumber": "10Y1532", "description": "Glove", "quantityShipped": 42.0, "unitOfMeasure": "PR",
               "salesOrder": None, "customerItemNumber": None}],
    "fieldIssues": [], "evidence": {}, "document": {"pageCount": 1}, "timings": {"totalSeconds": 1},
}
TWO_ITEMS = copy.deepcopy(ONE_ITEM)
TWO_ITEMS.update(shippedQuantity=None, unitOfMeasure="MG", items=[
    {"itemNumber": None, "description": "Silicon-28", "quantityShipped": 1000.0, "unitOfMeasure": "MG"},
    {"itemNumber": None, "description": "Silicon-29", "quantityShipped": 1000.0, "unitOfMeasure": "MG"}])


class EditedValuePOComparison(unittest.TestCase):
    def test_edited_item_quantity_drives_the_comparison(self):
        edited = copy.deepcopy(ONE_ITEM)
        edited["items"][0]["quantityShipped"] = "40"          # supplier corrected the line
        out = review.compare_with_po(edited, {"quantity": "42", "unit": "PR"})
        self.assertEqual(out["draft"]["shippedQuantity"], 40.0)  # total recomputed from the edit
        self.assertEqual(out["draft"]["quantitySummary"][0]["quantityShipped"], 40.0)  # summary follows the edit
        self.assertEqual(out["comparison"]["status"], "under_shipped")
        self.assertEqual(out["comparison"]["difference"], -2.0)

    def test_unit_spelling_is_normalised_before_comparing(self):
        edited = copy.deepcopy(ONE_ITEM)
        edited["items"][0]["unitOfMeasure"] = "pairs"
        self.assertEqual(review.compare_with_po(edited, {"quantity": 42, "unit": "pr"})["comparison"]["status"], "match")

    def test_different_units_are_never_compared(self):
        edited = copy.deepcopy(ONE_ITEM)
        edited["items"][0]["unitOfMeasure"] = "EA"
        c = review.compare_with_po(edited, {"quantity": 42, "unit": "PR"})["comparison"]
        self.assertEqual(c["status"], "cannot_compare")
        self.assertIsNone(c["difference"])

    def test_different_products_are_never_combined(self):
        out = review.compare_with_po(TWO_ITEMS, {"quantity": 2000, "unit": "MG"})
        self.assertIsNone(out["draft"]["shippedQuantity"])
        self.assertIn("different products", out["totals"]["message"])
        self.assertEqual(out["comparison"]["status"], "cannot_compare")

    def test_mixed_units_on_one_product_give_no_total(self):
        edited = copy.deepcopy(ONE_ITEM)
        edited["items"].append(dict(edited["items"][0], quantityShipped=3, unitOfMeasure="EA"))
        out = review.compare_with_po(edited, None)
        self.assertIsNone(out["draft"]["shippedQuantity"])
        self.assertIn("different units", out["totals"]["message"])

    def test_document_level_quantity_used_when_there_are_no_lines(self):
        edited = dict(ONE_ITEM, items=[], shippedQuantity="45", unitOfMeasure="PR")
        self.assertEqual(review.compare_with_po(edited, {"quantity": 42, "unit": "PR"})["comparison"]["status"],
                         "over_shipped")

    def test_invalid_po_entries_and_values_are_reported_not_compared(self):
        for po, field in [({"quantity": "abc", "unit": "PR"}, "poQuantity"), ({"quantity": "0", "unit": "PR"}, "poQuantity"),
                          ({"quantity": "5", "unit": ""}, "poUnit")]:
            out = review.compare_with_po(ONE_ITEM, po)
            self.assertIsNone(out["comparison"])
            self.assertIn(field, [e["field"] for e in out["errors"]])
        bad = copy.deepcopy(ONE_ITEM)
        bad["items"][0]["quantityShipped"] = "-3"
        out = review.compare_with_po(bad, {"quantity": 42, "unit": "PR"})
        self.assertIsNone(out["comparison"])
        self.assertEqual(out["errors"][0]["field"], "items[0].quantityShipped")

    def test_no_po_entered_means_no_comparison(self):
        out = review.compare_with_po(ONE_ITEM, {"quantity": "", "unit": "", "item": ""})
        self.assertFalse(out["poRequested"])
        self.assertIsNone(out["comparison"])


class FinishReview(unittest.TestCase):
    def setUp(self):
        self.original = copy.deepcopy(ONE_ITEM)
        self.original["fieldIssues"] = [{"field": "shipDate", "severity": "review", "code": "ambiguous_date",
                                         "message": "'2/7/2024' could be 7 Feb or 2 Jul."}]

    def test_unconfirmed_review_issue_blocks_finishing(self):
        res = review.finish_review(self.original, copy.deepcopy(self.original))
        self.assertFalse(res["ok"])
        self.assertIsNone(res["reviewed"])
        self.assertIn("Please confirm", res["errors"][0]["message"])

    def test_confirmed_review_produces_local_reviewed_json_with_corrections(self):
        edited = copy.deepcopy(self.original)
        edited["carrier"] = "UPS Ground"
        edited["trackingNumbers"] = ["1Z999AA10123456784", "  "]
        res = review.finish_review(self.original, edited, {"quantity": 42, "unit": "PR"}, ["shipDate:ambiguous_date"])
        self.assertTrue(res["ok"])
        rv = res["reviewed"]
        self.assertEqual(rv["status"], "reviewed_locally_not_submitted")
        self.assertIn("not posted to ERP", rv["review"]["note"])
        self.assertEqual(rv["trackingNumbers"], ["1Z999AA10123456784"])  # blank entry dropped
        fields = {c["field"] for c in rv["review"]["corrections"]}
        self.assertEqual(fields, {"carrier", "trackingNumbers"})
        self.assertEqual(rv["review"]["poComparison"]["status"], "match")
        self.assertNotIn("timings", rv)

    def test_invalid_date_blocks_and_clearing_a_value_is_allowed(self):
        edited = dict(copy.deepcopy(self.original), shipDate="2024-02-30")
        res = review.finish_review(self.original, edited, None, ["shipDate:ambiguous_date"])
        self.assertIn("shipDate", [e["field"] for e in res["errors"]])
        cleared = dict(copy.deepcopy(self.original), shipDate="", carrier=None)
        res = review.finish_review(self.original, cleared, None, ["shipDate:ambiguous_date"])
        self.assertTrue(res["ok"])
        self.assertIsNone(res["reviewed"]["shipDate"])
        self.assertIn("carrier", [w["field"] for w in res["warnings"]])


class ApiTests(unittest.TestCase):
    """Run the real HTTP server on a free port with a fake converter."""

    def start(self, convert, max_upload=1024 * 1024):
        self.app = ui_server.App(convert=convert, max_upload_bytes=max_upload)
        self.server = ui_server.make_server(self.app, "127.0.0.1", 0)
        self.base = f"http://127.0.0.1:{self.server.server_address[1]}"
        threading.Thread(target=self.server.serve_forever, daemon=True).start()
        self.addCleanup(self.server.server_close)
        self.addCleanup(self.server.shutdown)

    def call(self, path, data=None, ctype="application/json", method=None):
        body = json.dumps(data).encode() if isinstance(data, (dict, list)) else data
        req = urllib.request.Request(self.base + path, data=body, method=method or ("POST" if body is not None else "GET"),
                                     headers={"Content-Type": ctype} if body is not None else {})
        try:
            with urllib.request.urlopen(req, timeout=10) as r:
                return r.status, r.read()
        except urllib.error.HTTPError as e:
            return e.code, e.read()

    def upload(self, name, data):
        return self.call(f"/api/jobs?filename={name}", data, "application/pdf")

    def wait_job(self, job_id):
        for _ in range(100):
            status, body = self.call(f"/api/jobs/{job_id}")
            job = json.loads(body)
            if job["status"] != "running":
                return job
            time.sleep(0.05)
        self.fail("job did not finish")

    def test_upload_validation(self):
        self.start(lambda p: (dict(ONE_ITEM), None), max_upload=1000)
        self.assertEqual(self.upload("notes.txt", b"%PDF-1.4 x")[0], 415)
        self.assertEqual(self.upload("fake.pdf", b"PK\x03\x04 a zip")[0], 415)
        self.assertEqual(self.upload("empty.pdf", b"")[0], 400)
        status, body = self.upload("big.pdf", b"%PDF-" + b"x" * 2000)
        self.assertEqual(status, 413)
        self.assertIn("too large", json.loads(body)["error"]["message"])

    def test_successful_job_uses_safe_name_and_cleans_temp_files(self):
        seen = {}

        def fake_convert(pdf_path):
            seen["path"] = Path(pdf_path)
            seen["bytes"] = Path(pdf_path).read_bytes()
            return dict(ONE_ITEM), "# md"

        self.start(fake_convert)
        status, body = self.upload("..%5C..%5Cevil%3Cname%3E.pdf", b"%PDF-1.7 content")
        self.assertEqual(status, 202)
        job = self.wait_job(json.loads(body)["jobId"])
        self.assertEqual(job["status"], "done")
        self.assertEqual(job["result"]["sourceFile"], "evil_name_.pdf")
        self.assertEqual(seen["path"].name, "upload.pdf")          # user text never used as a path
        self.assertEqual(seen["bytes"], b"%PDF-1.7 content")
        self.assertFalse(seen["path"].parent.exists())             # temp folder removed

    def test_second_upload_while_busy_is_rejected(self):
        gate = threading.Event()
        self.start(lambda p: (gate.wait(5), (dict(ONE_ITEM), None))[1])
        status, body = self.upload("a.pdf", b"%PDF-1.4")
        self.assertEqual(status, 202)
        self.assertEqual(json.loads(self.call("/api/health")[1])["busy"], True)
        self.assertEqual(self.upload("b.pdf", b"%PDF-1.4")[0], 409)
        gate.set()
        self.assertEqual(self.wait_job(json.loads(body)["jobId"])["status"], "done")

    def test_conversion_error_is_friendly_and_has_no_traceback(self):
        failure = {"sourceFile": "upload.pdf", "status": "conversion_error", "conversionError": {
            "pdf": "C:/tmp/upload.pdf", "message": "failed", "attempts": [
                {"attempt": 1, "status": "crash", "reason": "access violation (0xC0000005)", "seconds": 3.1,
                 "stderrTail": ["Traceback (most recent call last):", "secret path C:/Users/x"]},
                {"attempt": 2, "status": "timeout", "reason": "exceeded 600 s", "seconds": 600.0,
                 "stderrTail": ["Traceback ..."]}]}}
        self.start(lambda p: (failure, None))
        job = self.wait_job(json.loads(self.upload("bad.pdf", b"%PDF-1.4")[1])["jobId"])
        self.assertEqual(job["status"], "failed")
        self.assertEqual(job["error"]["kind"], "conversion_error")
        self.assertEqual([a["outcome"] for a in job["error"]["attempts"]], ["crash", "timeout"])
        text = json.dumps(job)
        self.assertNotIn("Traceback", text)
        self.assertNotIn("stderrTail", text)
        self.assertNotIn("C:/tmp", text)

    def test_unexpected_server_error_is_reported_without_traceback(self):
        def boom(p):
            raise RuntimeError("internal detail that must not reach the browser")

        self.start(boom)
        from io import StringIO
        from contextlib import redirect_stderr
        with redirect_stderr(StringIO()):  # the server prints the traceback to its console only
            job = self.wait_job(json.loads(self.upload("x.pdf", b"%PDF-1.4")[1])["jobId"])
        self.assertEqual(job["error"]["kind"], "internal")
        self.assertNotIn("internal detail", json.dumps(job))

    def test_po_compare_and_review_endpoints_use_edited_values(self):
        self.start(lambda p: (dict(ONE_ITEM), None))
        edited = copy.deepcopy(ONE_ITEM)
        edited["items"][0]["quantityShipped"] = "44"
        status, body = self.call("/api/po-compare", {"draft": edited, "po": {"quantity": "42", "unit": "PR"}})
        out = json.loads(body)
        self.assertEqual((status, out["comparison"]["status"], out["shippedQuantity"]), (200, "over_shipped", 44.0))
        status, body = self.call("/api/review", {"original": ONE_ITEM, "draft": edited, "po": {"quantity": "42", "unit": "PR"},
                                                 "acknowledged": []})
        res = json.loads(body)
        self.assertTrue(res["ok"])
        self.assertEqual(res["reviewed"]["status"], "reviewed_locally_not_submitted")
        self.assertEqual(self.call("/api/review", b"not json")[0], 400)

    def test_static_files_and_path_safety(self):
        self.start(lambda p: (dict(ONE_ITEM), None))
        status, body = self.call("/")
        self.assertEqual(status, 200)
        self.assertIn(b"Packing-List Pre-Fill", body)
        self.assertEqual(self.call("/app.js")[0], 200)
        self.assertEqual(self.call("/../ui_server.py")[0], 404)
        self.assertEqual(self.call("/%2e%2e/review.py")[0], 404)
        self.assertEqual(self.call("/api/jobs/does-not-exist")[0], 404)

    def test_review_returns_summary_and_export_endpoint_serves_all_formats(self):
        self.start(lambda p: (dict(ONE_ITEM), None))
        edited = dict(copy.deepcopy(ONE_ITEM), carrier="FedEx Freight")
        res = json.loads(self.call("/api/review", {"original": ONE_ITEM, "draft": edited, "po": {}, "acknowledged": []})[1])
        self.assertEqual(dict(res["summary"]["shipment"])["Carrier"], "FedEx Freight")
        for fmt in ("txt", "pdf", "docx"):
            req = urllib.request.Request(f"{self.base}/api/export?format={fmt}", method="POST",
                                         data=json.dumps({"reviewed": res["reviewed"]}).encode(),
                                         headers={"Content-Type": "application/json"})
            with urllib.request.urlopen(req, timeout=10) as r:
                self.assertEqual(r.status, 200)
                self.assertIn(f'filename="x-reviewed-draft.{fmt}"', r.headers["Content-Disposition"])
                self.assertEqual(r.headers["Content-Type"], summary_export.EXPORT_FORMATS[fmt])
                self.assertIn("FedEx Freight", export_text(r.read(), fmt))
        self.assertEqual(self.call("/api/export?format=json", {"reviewed": res["reviewed"]})[0], 400)
        status, body = self.call("/api/export?format=pdf", {"reviewed": edited})  # not a finished review
        self.assertEqual(status, 400)
        self.assertIn("Finish the review first", json.loads(body)["error"]["message"])


def export_text(data, fmt):
    """Text content of an exported file (TXT decode, PDF via pypdfium2, DOCX via python-docx)."""
    if fmt == "txt":
        return data.decode("utf-8")
    if fmt == "pdf":
        import pypdfium2
        pdf = pypdfium2.PdfDocument(data)
        try:
            return "\n".join(pdf[i].get_textpage().get_text_range() for i in range(len(pdf)))
        finally:
            pdf.close()
    import io
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for row in t.rows for c in row.cells]
    parts += [p.text for s in doc.sections for p in s.header.paragraphs]
    return "\n".join(parts)


class SummaryExport(unittest.TestCase):
    def reviewed(self, **edits):
        original = copy.deepcopy(ONE_ITEM)
        original["items"][0].update(lots=[{"lotNumber": "2401"}, {"lotNumber": "9999"}], serialNumbers=[])
        original["fieldIssues"] = [{"field": "carrier", "severity": "review", "code": "ambiguous_carrier",
                                    "message": "Two carriers printed."}]
        edited = copy.deepcopy(original)
        edited.update(edits)
        res = review.finish_review(original, edited, {"quantity": 40, "unit": "PR"}, ["carrier:ambiguous_carrier"])
        self.assertTrue(res["ok"], res["errors"])
        return res["reviewed"]

    def test_summary_uses_final_edited_values_and_not_provided(self):
        s = summary_export.build_summary(self.reviewed(carrier="DHL Express", trackingNumbers=["1Z12345"]))
        shipment = dict(s["shipment"])
        self.assertEqual(shipment["Carrier"], "DHL Express")
        self.assertEqual(shipment["Ship date"], "2024-01-16 (16 January 2024)")
        self.assertEqual(shipment["Shipped quantity"], "42 PR")
        lists = dict(s["lists"])
        self.assertEqual(lists["Tracking numbers"], ["1Z12345"])
        self.assertEqual(lists["Serial numbers"], [])
        self.assertEqual(s["items"][0]["lots"], "2401")           # lot not in the final list is not shown
        self.assertEqual(s["items"][0]["serials"], "Not provided")
        po = dict(s["po"])
        self.assertEqual((po["Result"], po["Difference"]), ("Over-shipped", "+2 PR"))
        self.assertTrue(any(n.startswith("Corrected Carrier: UPS → DHL Express") for n in s["notes"]))
        self.assertTrue(any("Two carriers printed." in n for n in s["notes"]))
        self.assertEqual(s["banner"], "Reviewed draft — not submitted to ERP")

    def test_missing_values_say_not_provided_in_every_format(self):
        rv = self.reviewed(carrier=None, shipDate=None)
        s = summary_export.build_summary(rv)
        self.assertEqual(dict(s["shipment"])["Carrier"], "Not provided")
        self.assertEqual(dict(s["shipment"])["Ship date"], "Not provided")
        for fmt in ("txt", "pdf", "docx"):
            text = export_text(summary_export.export(rv, fmt)[0], fmt)
            self.assertIn("Carrier: Not provided" if fmt != "docx" else "Not provided", text)
            self.assertIn("Reviewed draft — not submitted to ERP", text)
            self.assertIn("10Y1532", text)
        pdf_text = export_text(summary_export.export(self.reviewed(carrier="DHL"), "pdf")[0], "pdf")
        self.assertIn("Corrected Carrier: UPS -> DHL", pdf_text)  # arrow is not in the PDF font

    def test_pdf_wraps_long_values_and_spans_pages(self):
        many = [f"1Z{n:016d}" for n in range(120)]
        rv = self.reviewed(trackingNumbers=many)
        data = summary_export.export(rv, "pdf")[0]
        import pypdfium2
        pdf = pypdfium2.PdfDocument(data)
        try:
            self.assertGreaterEqual(len(pdf), 1)
            text = "".join(pdf[i].get_textpage().get_text_range() for i in range(len(pdf))).replace("\r", "").replace("\n", "")
        finally:
            pdf.close()
        self.assertIn(many[-1], text)

    def test_export_rejects_unfinished_drafts_and_unknown_formats(self):
        with self.assertRaises(ValueError):
            summary_export.export(dict(ONE_ITEM), "pdf")
        with self.assertRaises(ValueError):
            summary_export.export(self.reviewed(), "json")


if __name__ == "__main__":
    unittest.main()
